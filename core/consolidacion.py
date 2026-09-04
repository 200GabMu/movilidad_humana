"""
Consolidación de los reportes mensuales (Excel) en una sola tabla.

Es la misma lógica del app.py original, con dos cambios de diseño:
  * No usa Streamlit: en vez de st.warning / st.error devuelve una lista de
    mensajes (nivel, texto) para que la interfaz decida cómo mostrarlos.
  * Recibe (nombre, bytes) en vez de objetos UploadedFile, así funciona igual
    desde la página web, desde un test o desde un script de consola.
"""

from __future__ import annotations

import io
import re
import warnings
from dataclasses import dataclass, field

import openpyxl
import pandas as pd

# ── Constantes de la plantilla ───────────────────────────────────────────────
HEADER_ROW_INDEX = 5
ROWS_TO_SKIP_AFTER_HEADER = 2
SHEET_PREFIX = "REPORTE MENSUAL"
MAX_DATA_COLUMNS = 79  # 78 en la plantilla sin KIT DE ALIMENTOS, 79 en la que sí lo trae
LAST_REAL_COLUMN = "OBTENCIÓN DE LA CÉDULA ECUATORIANA"

COL_HOJA = "Nombre de la Hoja"
COL_ARCHIVO = "Nombre del Archivo"
COLS_TRAZABILIDAD = [COL_HOJA, COL_ARCHIVO]


@dataclass
class Mensaje:
    """Un aviso generado durante el proceso. nivel: info | success | warning | error."""
    nivel: str
    texto: str


@dataclass
class ResultadoConsolidacion:
    df: pd.DataFrame | None
    mensajes: list[Mensaje] = field(default_factory=list)
    hojas_incluidas: int = 0
    filas: int = 0
    hojas_por_archivo: dict[str, list[str]] = field(default_factory=dict)


# ── Utilidades ───────────────────────────────────────────────────────────────

def get_valid_sheets(workbook: openpyxl.Workbook) -> list[str]:
    return [
        name for name in workbook.sheetnames
        if name.strip().upper().startswith(SHEET_PREFIX)
    ]


def extraer_año(nombre_archivo: str) -> str:
    """Extrae el año de 4 dígitos del nombre del archivo."""
    m = re.search(r"(20\d{2})", nombre_archivo)
    return m.group(1) if m else "DESCONOCIDO"


def normalizar_fecha(valor: str) -> str:
    """
    Convierte cualquier formato de fecha al formato YYYY/MM/DD.
    - '1995-06-20 00:00:00'  →  '1995/06/20'
    - '1995-06-20'           →  '1995/06/20'
    - '2012/09/11'           →  '2012/09/11'  (ya correcto)
    - Texto no reconocible lo deja como está.
    """
    if not valor or valor.strip() in ("", "nan", "Año/Mes/Día"):
        return valor
    v = valor.strip()
    v = re.sub(r"\s+\d{2}:\d{2}(:\d{2})?$", "", v)
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", v)
    if m:
        return f"{m.group(1)}/{m.group(2)}/{m.group(3)}"
    if re.match(r"^\d{4}/\d{2}/\d{2}$", v):
        return v
    return valor.strip()


def buscar_col(columnas, *palabras: str) -> str | None:
    """Primera columna cuyo nombre contenga TODAS las palabras (sin importar
    mayúsculas). Sirve para ubicar 'KIT DE ALIMENTOS' aunque esté escrito
    'KIT ALIMENTOS ' o con espacios de más."""
    for c in columnas:
        texto = str(c).upper()
        if all(p in texto for p in palabras):
            return c
    return None


# ── Lectura de una hoja ──────────────────────────────────────────────────────

def read_sheet(file_bytes: bytes, sheet_name: str) -> pd.DataFrame:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        df = pd.read_excel(
            io.BytesIO(file_bytes),
            sheet_name=sheet_name,
            header=HEADER_ROW_INDEX,
            dtype=str,
            engine="openpyxl",
        )
        df_fila5 = pd.read_excel(
            io.BytesIO(file_bytes),
            sheet_name=sheet_name,
            header=None,
            skiprows=4,
            nrows=1,
            dtype=str,
            engine="openpyxl",
        )

    # ── CORTE DURO INMEDIATO ──
    df = df.iloc[:, :MAX_DATA_COLUMNS]

    # ── CORRECCIÓN DE ENCABEZADOS COMBINADOS ──
    provisional_cols: list[str] = []
    for i, col in enumerate(df.columns):
        col_str = str(col).strip().replace("\n", " ")
        row7_val = (
            str(df.iloc[0, i]).strip().replace("\n", " ")
            if len(df) > 0 and not pd.isna(df.iloc[0, i]) else ""
        )

        if col_str.upper().startswith("ATENCIÓN HUMANITARIA") or col_str.startswith("Unnamed:"):
            if "KIT" in row7_val.upper() or "OTROS" in row7_val.upper():
                provisional_cols.append(row7_val)
                continue

        if col_str.upper() == "MADRE" or col_str.startswith("Unnamed:"):
            if "MADRE" in row7_val.upper():
                provisional_cols.append(row7_val)
                continue

        if col_str.upper() == "PADRE" or col_str.startswith("Unnamed:"):
            if "PADRE" in row7_val.upper():
                provisional_cols.append(row7_val)
                continue

        if col_str.upper().startswith("NNA SEPARADOS") or col_str.startswith("Unnamed:"):
            if any(k in row7_val.upper() for k in ["TUTOR", "CUIDADOR", "PARIENTE", "CARGO", "NNA"]):
                provisional_cols.append(row7_val)
                continue

        if col_str.startswith("Unnamed:") and i > 0 and provisional_cols and provisional_cols[-1] == "EDAD":
            provisional_cols[-1] = "EDAD - AÑOS"
            provisional_cols.append("EDAD - MESES")
            continue

        # Posiciones 42/43/44: el nombre se decide por el TEXTO de la fila 5,
        # porque si la plantilla trae KIT DE ALIMENTOS todo se corre una columna.
        if col_str.startswith("Unnamed:") and i in (42, 43, 44):
            if not df_fila5.empty and i < len(df_fila5.columns):
                val_fila5 = str(df_fila5.iloc[0, i]).strip().replace("\n", " ")
                if val_fila5 and not val_fila5.startswith("nan"):
                    if "VULNERABILIDAD" in val_fila5.upper():
                        provisional_cols.append("OTRO TIPO DE VULNERABILIDAD")
                    else:
                        provisional_cols.append("SÓLO PARA CASO DE RESPUESTA OTRO")
                    continue

        if col_str.startswith("Unnamed:"):
            col_str = f"Columna_{i}"

        provisional_cols.append(col_str)

    # ── ANTI-DUPLICADOS ──
    final_cols: list[str] = []
    seen: dict[str, int] = {}
    for c in provisional_cols:
        if c in seen:
            seen[c] += 1
            final_cols.append(f"{c}__{seen[c]}")
        else:
            seen[c] = 0
            final_cols.append(c)

    # ── CORTE POR NOMBRE DE COLUMNA FINAL ──
    cut_idx = None
    for idx in range(len(final_cols) - 1, -1, -1):
        clean = re.sub(r"__\d+$", "", final_cols[idx])
        if LAST_REAL_COLUMN in clean.upper():
            cut_idx = idx
            break

    if cut_idx is not None:
        final_cols = final_cols[: cut_idx + 1]
        df = df.iloc[:, : cut_idx + 1]
    else:
        last_real_idx = 0
        for idx in range(len(final_cols) - 1, -1, -1):
            col_name = final_cols[idx]
            if not col_name.startswith("Columna_") and not col_name.startswith("Unnamed:"):
                last_real_idx = idx
                break
        final_cols = final_cols[: last_real_idx + 1]
        df = df.iloc[:, : last_real_idx + 1]

    df.columns = final_cols
    df = df.iloc[ROWS_TO_SKIP_AFTER_HEADER:].reset_index(drop=True)

    # ── ELIMINACIÓN DE FILAS FANTASMAS ──
    df = df.replace(r"^\s*$", None, regex=True)
    df = df.dropna(how="all")

    if not df.empty:
        primer_col = df.columns[0]
        es_nulo = df[primer_col].isna()
        es_basura = (
            df[primer_col].astype(str).str.strip().str.lower()
            .isin(["nan", "none", "<na>", "", "n/d"])
        )
        df = df[~(es_nulo | es_basura)].reset_index(drop=True)

    # Filas que solo traen "N°" (autonumeración de plantilla) y nada más
    if not df.empty:
        cols_a_revisar = [c for c in df.columns if c != df.columns[0]]
        if cols_a_revisar:
            fila_vacia = df[cols_a_revisar].isna().all(axis=1)
            df = df[~fila_vacia].reset_index(drop=True)

    # ── NORMALIZAR FECHA DE NACIMIENTO ──
    fecha_col = next(
        (c for c in df.columns if "FECHA" in str(c).upper() and "NACIMIENTO" in str(c).upper()),
        None,
    )
    if fecha_col:
        df[fecha_col] = df[fecha_col].apply(
            lambda x: normalizar_fecha(str(x)) if pd.notna(x) else x
        )

    return df


def forzar_kit_alimentos(df: pd.DataFrame) -> pd.DataFrame:
    """Si algún archivo trae KIT DE ALIMENTOS, la deja justo después de
    KIT DE SALUD (orden de plantilla: ASEO · SALUD · ALIMENTOS · ESCOLAR · OTROS).
    No se inventa cuando ningún archivo la tiene: el consolidado de referencia
    no la incluye."""
    df = df.copy()
    col_alimentos = buscar_col(df.columns, "KIT", "ALIMENTO")
    if col_alimentos is None:
        return df

    col_salud = buscar_col(df.columns, "KIT", "SALUD")
    if col_salud is None:
        return df

    cols = list(df.columns)
    cols.remove(col_alimentos)
    cols.insert(cols.index(col_salud) + 1, col_alimentos)
    return df[cols]


# ── Alineación por posición con la plantilla ─────────────────────────────────
# La plantilla mensual SIEMPRE tiene las mismas 78 columnas en el mismo orden;
# lo que cambia entre años son los nombres ('Edad años' vs 'EDAD - AÑOS',
# 'TELÉFONO' en vez de 'Teléfono de contacto de la MADRE', un encabezado
# vacío, una columna en blanco de más…). Emparejar por NOMBRE perdía datos y
# corría las columnas repetidas "SÓLO PARA CASO DE RESPUESTA OTRO"; emparejar
# por POSICIÓN, comprobando unas columnas ancla, es exacto.
_ANCLAS = {0: "N°", 1: "ZONA", 2: "PROVINCIA", 4: "CIUDAD", 6: "APELLIDOS", 7: "FECHA DE NACIMIENTO",
           11: "SEXO", 13: "NACIONALIDAD", 17: "DOCUMENTOS DE VIAJE", 23: "ATENCIÓN EMERGENTE",
           30: "CONDICIÓN SOLO PARA NNA", 45: "TIENE DISCAPACIDAD", 50: "ATENCIÓN DE TRABAJO SOCIAL",
           77: "OBTENCIÓN DE LA CÉDULA ECUATORIANA"}


def _clave_pos(texto) -> str:
    import unicodedata
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", t).strip().upper()


def _plantilla_datos(con_kit_alimentos: bool) -> list[str]:
    from .plantilla import COLUMNAS_SALIDA  # import tardío: evita ciclo
    cols = [c for c in COLUMNAS_SALIDA if c not in (COL_HOJA, COL_ARCHIVO, "MES", "AÑO")]
    if con_kit_alimentos:
        i = cols.index("KIT DE SALUD") + 1
        cols = cols[:i] + ["KIT DE ALIMENTOS"] + cols[i:]
    return cols


def alinear_por_posicion(df: pd.DataFrame) -> tuple[pd.DataFrame, str | None]:
    """Renombra las columnas de una hoja según su POSICIÓN en la plantilla.
    Devuelve (df, aviso). Si la hoja no encaja con la plantilla (anclas
    distintas) se deja como está y se explica en el aviso."""
    df = df.copy()
    # columnas sin encabezado y completamente vacías (una columna en blanco de más)
    sobrantes = [c for c in df.columns if str(c).startswith(("Columna_", "Unnamed:")) and df[c].isna().all()]
    while sobrantes and len(df.columns) > 78:
        df = df.drop(columns=sobrantes.pop(0))
    cols = list(df.columns)
    con_kit = len(cols) == 79 and "ALIMENT" in _clave_pos(cols[29])
    plantilla = _plantilla_datos(con_kit)
    if len(cols) != len(plantilla):
        return df, f"tiene {len(cols)} columnas de datos y la plantilla {len(plantilla)}; se emparejó por nombre"
    fallos = [f"col {i + 1}: '{cols[i]}' ≠ '{nombre}'" for i, nombre in _ANCLAS.items()
              if not _clave_pos(cols[i]).startswith(_clave_pos(nombre))]
    if len(fallos) > 2:
        return df, "no coincide con la plantilla (" + "; ".join(fallos[:3]) + "); se emparejó por nombre"
    df.columns = plantilla
    return df, None


def add_traceability_columns(df: pd.DataFrame, sheet_name: str, file_name: str) -> pd.DataFrame:
    df = df.copy()
    df[COL_HOJA] = sheet_name
    df[COL_ARCHIVO] = file_name
    return df


# ── Procesar un archivo ──────────────────────────────────────────────────────

def process_file(file_bytes: bytes, file_name: str) -> tuple[pd.DataFrame | None, list[str], list[str], list[Mensaje]]:
    """Devuelve (df, hojas procesadas, hojas omitidas, mensajes)."""
    mensajes: list[Mensaje] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # archivo corrupto, no es xlsx, etc.
        mensajes.append(Mensaje("error", f"No se pudo abrir **{file_name}**: {exc}"))
        return None, [], [], mensajes

    all_sheets = wb.sheetnames
    valid_sheets = get_valid_sheets(wb)
    skipped = [s for s in all_sheets if s not in valid_sheets]

    if not valid_sheets:
        mensajes.append(Mensaje(
            "warning",
            f"**{file_name}** no contiene hojas que empiecen por '{SHEET_PREFIX}'. Archivo ignorado.",
        ))
        return None, [], skipped, mensajes

    frames: list[pd.DataFrame] = []
    processed: list[str] = []

    for sheet in valid_sheets:
        try:
            df_sheet = read_sheet(file_bytes, sheet)
            df_sheet, aviso = alinear_por_posicion(df_sheet)
            if aviso:
                mensajes.append(Mensaje("warning", f"Hoja **{sheet}** de **{file_name}**: {aviso}."))
            df_sheet = add_traceability_columns(df_sheet, sheet, file_name)
            frames.append(df_sheet)
            processed.append(sheet)
        except Exception as exc:
            mensajes.append(Mensaje("warning", f"Error leyendo hoja **{sheet}** en **{file_name}**: {exc}"))

    if not frames:
        return None, processed, skipped, mensajes

    return pd.concat(frames, ignore_index=True), processed, skipped, mensajes


# ── Unir todo ────────────────────────────────────────────────────────────────

def clave_hoja(nombre_archivo: str, nombre_hoja: str, n_filas: int) -> tuple:
    """Clave para detectar la misma hoja subida dos veces: (año, hoja, filas)."""
    return (extraer_año(nombre_archivo), nombre_hoja.strip().upper(), int(n_filas))


def claves_de(df: pd.DataFrame) -> set[tuple]:
    """Claves de todas las hojas que ya están dentro de un consolidado."""
    if df is None or df.empty or COL_HOJA not in df.columns:
        return set()
    conteo = df.groupby([COL_ARCHIVO, COL_HOJA]).size()
    return {clave_hoja(a, h, n) for (a, h), n in conteo.items()}


def alinear_columnas(frames: list[pd.DataFrame]) -> pd.DataFrame:
    """Concatena hojas con distintas columnas sin perder ninguna.

    La referencia es la hoja MÁS ANCHA; las columnas que solo existen en
    otras hojas se agregan al final de los datos (antes de la trazabilidad).
    """
    ref_cols = max(frames, key=lambda f: len(f.columns)).columns.tolist()
    cut_ref = None
    for idx_c, col in enumerate(ref_cols):
        clean = re.sub(r"__\d+$", "", col)
        if LAST_REAL_COLUMN in clean.upper():
            cut_ref = idx_c
            break

    if cut_ref is not None:
        data_cols = ref_cols[: cut_ref + 1]
    else:
        data_cols = [c for c in ref_cols if c not in COLS_TRAZABILIDAD]

    keep_cols = data_cols + list(COLS_TRAZABILIDAD)
    for frame in frames:
        for c in frame.columns:
            if c not in keep_cols:
                keep_cols.insert(len(keep_cols) - 2, c)

    safe_frames = [frame.reindex(columns=keep_cols) for frame in frames]
    df_final = pd.concat(safe_frames, ignore_index=True)
    return forzar_kit_alimentos(df_final)


def unify_files(
    archivos: list[tuple[str, bytes]],
    claves_existentes: set[tuple] | None = None,
    progreso=None,
) -> ResultadoConsolidacion:
    """
    Consolida una lista de archivos (nombre, bytes).

    claves_existentes: hojas que YA están en el consolidado guardado; se
        omiten para no duplicar cuando el usuario sube un archivo repetido.
    progreso: callback opcional progreso(fraccion, texto) para la barra de la UI.
    """
    res = ResultadoConsolidacion(df=None)
    claves_vistas: set[tuple] = set(claves_existentes or set())
    all_frames: list[pd.DataFrame] = []

    if not archivos:
        res.mensajes.append(Mensaje("error", "No se recibió ningún archivo."))
        return res

    for idx, (file_name, file_bytes) in enumerate(archivos):
        res.mensajes.append(Mensaje("info", f"📄 Procesando `{file_name}`"))
        df_file, processed, skipped, msgs = process_file(file_bytes, file_name)
        res.mensajes.extend(msgs)

        if df_file is not None:
            for sheet_name in processed:
                df_hoja = df_file[df_file[COL_HOJA] == sheet_name]
                n_filas = len(df_hoja)
                clave = clave_hoja(file_name, sheet_name, n_filas)

                if clave in claves_vistas:
                    res.mensajes.append(Mensaje(
                        "warning",
                        f"Hoja duplicada ignorada: *{sheet_name}* "
                        f"(año={clave[0]}, filas={n_filas}) — ya existe en el consolidado o en otro archivo.",
                    ))
                else:
                    claves_vistas.add(clave)
                    all_frames.append(df_hoja)
                    res.hojas_incluidas += 1
                    res.filas += n_filas
                    res.hojas_por_archivo.setdefault(file_name, []).append(sheet_name)
                    res.mensajes.append(Mensaje("success", f"Hoja incluida: *{sheet_name}* ({n_filas} filas)"))

        for sheet in skipped:
            res.mensajes.append(Mensaje("info", f"⏭️ Hoja ignorada (no cumple prefijo): *{sheet}*"))

        if progreso:
            progreso((idx + 1) / len(archivos), f"Procesado {idx + 1} de {len(archivos)} archivos…")

    if not all_frames:
        res.mensajes.append(Mensaje("error", "Ningún archivo contenía hojas nuevas para consolidar."))
        return res

    res.df = alinear_columnas(all_frames)
    return res


def fusionar(df_existente: pd.DataFrame | None, df_nuevo: pd.DataFrame) -> pd.DataFrame:
    """Une el consolidado guardado con las hojas recién cargadas (modo incremental)."""
    if df_existente is None or df_existente.empty:
        return df_nuevo
    return alinear_columnas([df_existente, df_nuevo])

"""Genera Excel sintéticos con la MISMA estructura de la plantilla (encabezado
en fila 6, datos desde la fila 9) y las columnas reales del formulario, para
probar la app y el panel sin datos personales.   python tests/crear_excel_prueba.py"""

import random
from pathlib import Path

from openpyxl import Workbook

random.seed(7)
SALIDA = Path(__file__).resolve().parent / "datos_prueba"

SI_NO = ["SI", "NO", "NO", "NO", "No aplica"]
COLUMNAS = [
    "N°", "ZONA", "PROVINCIA", "DISTRITO", "CIUDAD", "NOMBRES", "APELLIDOS", "FECHA DE NACIMIENTO",
    "EDAD - AÑOS", "EDAD - MESES", "RANGO DE EDAD", "SEXO", "GÉNERO", "NACIONALIDAD", "ETNIA",
    "SÓLO PARA CASO DE RESPUESTA OTRO", "DIRECCIÓN DOMICILIAR", "DOCUMENTOS DE VIAJE",
    "NÚMERO DEL DOCUMENTO DE VIAJE", "SITUACIÓN DE MOVILIDAD", "FORMA DE INGRESO AL ECUADOR",
    "Motivo de salida del país de origen", "ATENCIÓN EMERGENTE",
    "INDICAR QUE TIPO DE ATENCIÓN EMERGENTE ACTUAL RECIBE", "KIT DE ASEO", "KIT DE SALUD",
    "KIT ESCOLAR", "OTROS (Se debe especificar que otro tipo de KIT humanitario fue entregado)",
    "CONDICIÓN SOLO PARA NNA", "CON QUIEN SE ENCUENTRA EL NNA", "OTRO TIPO DE VULNERABILIDAD",
    "TIENE ENFERMEDAD CATASTRÓFICA", "TIENE DISCAPACIDAD", "Tipo de discapacidad",
    "EMBARAZO  (Sólo para mujeres)", "ACTUALMENTE ESTÁ ESTUDIANDO", "NIVEL DE ESCOLARIDAD / INSTRUCCIÓN",
    "ATENCIÓN DE TRABAJO SOCIAL", "ATENCIÓN PSICOLÓGICA", "ATENCIÓN LEGAL", "SALUD", "EDUCACIÓN",
    "PARTICIPACIÓN A TALLERES DE CAPACITACIÓN", "PARTICIPACIÓN A TALLERES DE SENSIBILIZACIÓN",
    "PARTICIPACIÓN EN ENCUENTROS COMUNITARIOS",
    "PARTICIPACIÓN EN TALLERES Y ESPACIOS LÚDICO DE RECREACIÓN DE NNA", "INSERCIÓN A REDES COMUNITARIAS",
    "SITUACIÓN  MIGRATORIA", "REGISTRO MIGRATORIO  -MDI-", "VISA VIRTE -MREMH-",
    "OBTENCIÓN DE LA CÉDULA ECUATORIANA -REGISTRO CIVIL-", "NOTAS INTERNAS (fuera de matriz)",
]

VALORES = {
    "ZONA": ["ZONA 6"] * 20 + ["ZONA 7", "ZONA 8"],
    "PROVINCIA": ["AZUAY"], "DISTRITO": ["01D01"], "CIUDAD": ["CUENCA", "Cuenca"],
    "NOMBRES": ["MARÍA", "JOSÉ", "Ana", "luis", "CARMEN", "PEDRO", "Rosa", "DIEGO", "LUISA", "JUAN"],
    "APELLIDOS": ["PÉREZ", "GÓMEZ", "Rodríguez", "LÓPEZ", "Martínez", "SILVA", "TORRES", "RAMOS"],
    "RANGO DE EDAD": ["ADULTO"] * 5 + ["NN"] * 4 + ["ADOLESCENTE"] + ["ADULTO MAYOR"],
    "SEXO": ["MUJER", "MUJER", "HOMBRE", "Seleccione"],
    "GÉNERO": ["FEMENINO", "FEMENINO", "MASCULINO", "LGTBIQ+"],
    "NACIONALIDAD": ["VENEZOLANA", "Venezolano", "VENEZOLANA", "ECUATORIANO", "COLOMBIANA", "PERUANA", "Seleccione"],
    "ETNIA": ["MESTIZO", "AFRODESCENDIENTE", "No aplica"],
    "SITUACIÓN DE MOVILIDAD": ["VOCACION DE PERMANENCIA"] * 9 + ["EN TRANSITO"],
    "FORMA DE INGRESO AL ECUADOR": ["POR PASO IRREGULAR"] * 5 + ["POR PASO REGULAR", "No aplica"],
    "Motivo de salida del país de origen": [
        "BÚSQUEDA DE TRABAJO Y MEJORAS DE LA CALIDAD DE VIDA", "FALTA DE ACCESO A MEDICINA, ALIMENTACIÓN Y SERVICIOS BÁSICOS",
        "ESTUDIOS", "SITUACIÓN SOCIOECONÓMICA", "AMENAZAS CONTRA LA VIDA, SEGURIDAD, LIBERTAD E INTEGRIDAD", "No aplica",
    ],
    "INDICAR QUE TIPO DE ATENCIÓN EMERGENTE ACTUAL RECIBE": ["No aplica"] * 8 + ["ALOJAMIENTO TEMPORAL", "ALIMENTOS"],
    "OTROS (Se debe especificar que otro tipo de KIT humanitario fue entregado)": ["No aplica", "ALIMENTOS"],
    "CONDICIÓN SOLO PARA NNA": ["NNA -ACOMPAÑADO", "No aplica", "NNA - SEPARADO", "NNA - NO ACOMPAÑADO"],
    "CON QUIEN SE ENCUENTRA EL NNA": ["No aplica", "SOLO CON LA MADRE", "CON AMBOS", "NINGUNO", "SOLO CON EL PADRE"],
    "OTRO TIPO DE VULNERABILIDAD": ["OTRO TIPO DE VULNERABILIDAD", "No aplica", "MENDICIDAD", "TRABAJO INFANTIL"],
    "Tipo de discapacidad": ["No aplica"] * 8 + ["FISICA", "INTELECTUAL", "MULTIPLE"],
    "NIVEL DE ESCOLARIDAD / INSTRUCCIÓN": ["BASICA", "BACHILLERATO", "SUPERIOR", "ANALFABETO", "INICIAL", "No aplica"],
    "SITUACIÓN  MIGRATORIA": ["IRREGULAR", "IRREGULAR", "REGULAR", "No aplica"],
    "REGISTRO MIGRATORIO  -MDI-": ["NO", "NO", "SI", "EN TRAMITE"],
    "VISA VIRTE -MREMH-": ["NO", "NO", "SI", "EN TRAMITE"],
    "OBTENCIÓN DE LA CÉDULA ECUATORIANA -REGISTRO CIVIL-": ["NO", "NO", "SI", "EN TRAMITE"],
    "NOTAS INTERNAS (fuera de matriz)": ["no debería entrar"],
}


def fila_persona(k: int, año: int) -> list:
    edad = random.choice([None, random.randint(0, 75)])
    datos = []
    for col in COLUMNAS:
        if col == "N°":
            datos.append(k)
        elif col == "FECHA DE NACIMIENTO":
            datos.append(f"{año - (edad or 30)}-0{random.randint(1, 9)}-1{random.randint(0, 9)} 00:00:00")
        elif col == "EDAD - AÑOS":
            datos.append(edad)
        elif col == "EDAD - MESES":
            datos.append(random.choice([None, 3]))
        elif col in ("SÓLO PARA CASO DE RESPUESTA OTRO", "DIRECCIÓN DOMICILIAR", "DOCUMENTOS DE VIAJE",
                     "NÚMERO DEL DOCUMENTO DE VIAJE"):
            datos.append(random.choice([None, "No aplica"]))
        elif col in VALORES:
            datos.append(random.choice(VALORES[col]))
        else:  # casillas SI / NO
            datos.append(random.choice(SI_NO))
    return datos


def hoja(wb, titulo, n, año):
    ws = wb.create_sheet(titulo)
    ws["A1"] = "FUNDACIÓN MENSAJEROS DE LA PAZ"
    ws["A2"] = "SERVICIO DE MOVILIDAD HUMANA"
    ws["A5"] = "DATOS PERSONALES"
    for i, h in enumerate(COLUMNAS, start=1):
        ws.cell(row=6, column=i, value=h)
    ws.cell(row=8, column=8, value="Año/Mes/Día")   # fila de ejemplo de la plantilla
    fila = 9
    for k in range(1, n + 1):
        for i, v in enumerate(fila_persona(k, año), start=1):
            ws.cell(row=fila, column=i, value=v)
        fila += 1
    for k in range(3):  # filas fantasma con solo N°
        ws.cell(row=fila + k, column=1, value=n + k + 1)
    return ws


def crear(nombre, año, hojas):
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("INSTRUCCIONES")["A1"] = "hoja que se debe ignorar"
    for titulo, n in hojas:
        hoja(wb, titulo, n, año)
    SALIDA.mkdir(parents=True, exist_ok=True)
    ruta = SALIDA / nombre
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    print(crear("REPORTE MENSUAL 2025.xlsx", 2025, [("REPORTE MENSUAL OCTUBRE", 40), ("REPORTE MENSUAL NOVIEMBRE", 35)]))
    print(crear("REPORTE MENSUAL 2026.xlsx", 2026, [("REPORTE MENSUAL ENERO", 50), ("REPORTE MENSUAL FEBRERO", 45)]))
    print(crear("REPORTE MENSUAL 2026 marzo.xlsx", 2026, [("REPORTE MENSUAL MARZO", 30)]))
